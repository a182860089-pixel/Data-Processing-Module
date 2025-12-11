"""
Office 文档转 PDF 转换器
支持 DOCX、PPTX、XLSX 等格式转换为 PDF
使用 LibreOffice 后端实现高质量转换
"""
import logging
from typing import Dict, Any, List
from pathlib import Path
import subprocess
import tempfile
from app.core.base.converter import BaseConverter, ConversionResult
from app.exceptions.converter_exceptions import ConversionFailedException

logger = logging.getLogger(__name__)


class OfficeConverter(BaseConverter):
    """Office 文档转 PDF 转换器"""
    
    def __init__(self):
        """初始化转换器"""
        self.supported_extensions = ['.docx', '.doc', '.pptx', '.ppt', '.xlsx', '.xls']
    
    async def convert(
        self,
        file_path: str,
        options: Dict[str, Any]
    ) -> ConversionResult:
        """
        使用 LibreOffice 将 Office 文档转换为 PDF
        
        Args:
            file_path: 源文件路径
            options: 转换选项
            
        Returns:
            ConversionResult: 转换结果
        """
        logger.info(f"Starting Office conversion with LibreOffice: {file_path}")
        
        try:
            ext = Path(file_path).suffix.lower()
            
            # 验证文件类型
            if ext in ['.docx', '.doc']:
                doc_type = 'docx'
            elif ext in ['.pptx', '.ppt']:
                doc_type = 'pptx'
            elif ext in ['.xlsx', '.xls']:
                doc_type = 'xlsx'
            else:
                raise ConversionFailedException(
                    message=f"不支持的 Office 文件类型: {ext}"
                )
            
            # 使用 LibreOffice 转换
            pdf_content = self._convert_with_libreoffice(file_path)
            
            logger.info(f"Office conversion completed: {len(pdf_content)} bytes")
            
            result = ConversionResult(
                pdf_content=pdf_content,
                metadata={
                    'format': 'pdf',
                    'source_type': doc_type,
                    'file_size': len(pdf_content),
                    'converter': 'LibreOffice',
                },
                status='success',
                output_type='pdf'
            )
            
            return result
            
        except Exception as e:
            logger.error(f"Office conversion failed: {str(e)}")
            raise ConversionFailedException(
                message="Office 文档转换失败",
                details=str(e)
            )
    
    def _convert_with_libreoffice(self, file_path: str) -> bytes:
        """
        使用 LibreOffice 命令行转换 Office 文档为 PDF
        
        Args:
            file_path: 源文件路径
            
        Returns:
            bytes: PDF 内容
        """
        try:
            # 使用临时目录为输出
            with tempfile.TemporaryDirectory() as tmpdir:
                logger.info(f"Converting with LibreOffice: {file_path} -> {tmpdir}")
                
                # LibreOffice 位置
                libreoffice_path = r"C:\Program Files\LibreOffice\program\soffice.exe"
                
                # 执行 LibreOffice 命令: --headless 后台运行 --convert-to pdf
                cmd = [
                    libreoffice_path,
                    "--headless",
                    "--convert-to", "pdf",
                    "--outdir", tmpdir,
                    file_path
                ]
                
                logger.info(f"Executing command: {' '.join(cmd)}")
                
                # 使用 subprocess 执行转换
                result = subprocess.run(
                    cmd,
                    capture_output=True,
                    timeout=60,
                    text=True
                )
                
                if result.returncode != 0:
                    logger.error(f"LibreOffice error: {result.stderr}")
                    raise Exception(f"LibreOffice conversion failed: {result.stderr}")
                
                # 找到生成的 PDF 文件
                pdf_files = list(Path(tmpdir).glob("*.pdf"))
                if not pdf_files:
                    raise Exception("PDF file not generated")
                
                pdf_path = pdf_files[0]
                logger.info(f"PDF generated: {pdf_path}")
                
                # 读取 PDF 文件内容
                with open(pdf_path, 'rb') as f:
                    pdf_content = f.read()
                
                return pdf_content
                
        except Exception as e:
            logger.error(f"LibreOffice conversion failed: {str(e)}")
            raise ConversionFailedException(
                message="Office 文档转换失败",
                details=str(e)
            )
    
    def validate(self, file_path: str) -> bool:
        """验证是否为有效的 Office 文件"""
        try:
            ext = Path(file_path).suffix.lower()
            
            if ext in ['.docx', '.doc']:
                from docx import Document
                doc = Document(file_path)
                return True
            elif ext in ['.pptx', '.ppt']:
                from pptx import Presentation
                prs = Presentation(file_path)
                return True
            elif ext in ['.xlsx', '.xls']:
                from openpyxl import load_workbook
                wb = load_workbook(file_path)
                return True
            
            return False
        except Exception as e:
            logger.error(f"Office file validation failed: {str(e)}")
            return False
    
    def get_supported_extensions(self) -> List[str]:
        """返回支持的文件扩展名列表"""
        return self.supported_extensions

